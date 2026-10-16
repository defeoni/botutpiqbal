"""Read the official item workbook and produce a deterministic local CSV.

Usage: python alat/siapkan-master-item.py "C:/path/PT.UTP - Master Item.xlsx"
Requires openpyxl for reading only. The input workbook is never saved or edited.
"""

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path


FIELDS = ["code", "name", "brand", "brandCandidate", "brandStatus", "model",
          "unit", "group", "family"]
SHEET = "Mast, Items (Stock...)"
SOURCE_HEADERS = {"code": "Kode Item", "name": "Nama Item", "brand": "Merk",
                  "unit": "Sat", "group": "Kelompok", "family": "Family"}
GENERIC = {
    "", "-", "?", "n/a", "na", "none", "unknown", "tanpa merk", "tanpa merek",
    "no brand", "non brand", "generic", "generik", "lokal", "local", "import",
    "china", "standar", "standard", "setara", "equivalent", "compare", "modul",
    "riser", "silver", "gold", "hitam", "putih", "merah", "hijau", "biru",
    "kuning", "abu", "abu abu", "abu-abu", "coklat", "orange", "black", "white",
    "red", "green", "blue", "grey", "gray", "yellow", "brown", "clear", "las",
    "pvc", "upvc", "cpvc", "hdpe", "ppr", "ss304", "ss316", "swr", "bs",
}
DESCRIPTOR_WORDS = {
    "mm", "cm", "inch", "watt", "volt", "kw", "kva", "btu", "psi", "bar",
    "sch", "schedule", "class", "pn", "dn", "diameter", "dia", "tebal",
    "besi", "baja", "plat", "papan", "triplek", "kayu", "karet", "tembaga",
    "galvanis", "galvanized", "galvanised", "stainless", "steel", "aluminium",
    "aluminum", "indoor", "outdoor", "unit", "wall", "mounting", "standing",
    "drat", "dalam", "luar", "ulir", "baut", "hole", "holes", "type", "tipe",
    "name", "panel", "flange", "flanged", "socket", "thread", "threaded",
    "pressure", "flat", "round", "plate", "hollow", "color", "warna", "jenis",
}


def text(value):
    """Preserve source strings, including deliberate spaces and spelling."""
    return "" if value is None else str(value)


def normalize(value):
    return " ".join(text(value).strip().casefold().split())


def is_generic(value):
    key = normalize(value)
    words = set(re.findall(r"[a-z]+", key))
    return key in GENERIC or bool(words & DESCRIPTOR_WORDS)


def plausible_brand(value, structured):
    key = normalize(value)
    if key in structured:
        return True
    if is_generic(value) or not 2 <= len(key) <= 45 or len(key.split()) > 4:
        return False
    # Capitalization alone is never evidence. Reject dimensions/specifications.
    if not re.fullmatch(r"[\w &.'/-]+", key, re.UNICODE):
        return False
    if not any(c.isalpha() for c in key):
        return False
    if re.search(r"\d\s*[x×/.,]\s*\d|\b\d+\s*(?:mm|cm|m|w|v|a|kg|hz)\b", key):
        return False
    return not key[0].isdigit() or key == "3m"


def name_parts(name, structured):
    """Find the final comma segment preceding a model delimiter.

    Prefer a pipe, since a hash can also introduce a thickness or dimension.
    A final comma segment without a delimiter may suggest a brand, never a model.
    """
    choices = []
    for delimiter in ("|", "#"):
        for match in re.finditer(re.escape(delimiter), name):
            prefix = name[:match.start()]
            if "," not in prefix:
                continue
            fragment = prefix.rsplit(",", 1)[1].strip()
            suffix = name[match.end():].strip()
            choices.append((fragment, suffix, delimiter))
    for fragment, suffix, delimiter in choices:
        if plausible_brand(fragment, structured):
            return fragment, suffix, delimiter
    # Keep a clear pipe model even when its neighboring descriptor is a color.
    for fragment, suffix, delimiter in choices:
        if delimiter == "|":
            return fragment, suffix, delimiter
    if choices:
        return choices[0]
    return (name.rsplit(",", 1)[1].strip() if "," in name else ""), "", ""


def parse_item(name, raw_brand, structured, repeated):
    fragment, suffix, delimiter = name_parts(name, structured)
    candidate = fragment if plausible_brand(fragment, structured) else ""
    raw_key = normalize(raw_brand)
    if raw_key:
        status = "unknown" if raw_key in GENERIC else "explicit"
    elif candidate:
        status = "candidate" if normalize(candidate) in structured | repeated else "unknown"
    else:
        status = "unknown"
    # Do not present size expressions introduced by # as a model number.
    dimension = bool(re.fullmatch(r"(?=\s*\d)[\d\s.,x×/-]+(?:mm|cm|m|inch|in|kg|w|v)?", suffix, re.I))
    model = ""
    evidence = "none"
    if suffix and not dimension and (delimiter == "|" or (delimiter == "#" and candidate)):
        model = suffix
        evidence = "pipe_delimiter" if delimiter == "|" else "hash_after_brand_candidate"
    return {"brandCandidate": candidate, "brandStatus": status, "model": model}, evidence


def code_text(cell):
    """Keep text identifiers intact; honor a numeric zero-padding format."""
    value = cell.value
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        raise ValueError(f"Boolean item code at {cell.coordinate}")
    if isinstance(value, (int, float)):
        if not float(value).is_integer():
            raise ValueError(f"Non-integral numeric item code at {cell.coordinate}")
        number = str(int(value))
        fmt = cell.number_format.split(";", 1)[0]
        return number.zfill(len(fmt)) if re.fullmatch(r"0+", fmt) else number
    raise ValueError(f"Unsupported item code at {cell.coordinate}")


def read_items(source, sheet):
    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        worksheet = workbook[sheet]
        header_row = None
        columns = None
        for cells in worksheet.iter_rows(max_row=20):
            headings = {normalize(c.value): c.column - 1 for c in cells if c.value is not None}
            if all(normalize(label) in headings for label in SOURCE_HEADERS.values()):
                header_row = cells[0].row
                columns = {key: headings[normalize(label)] for key, label in SOURCE_HEADERS.items()}
                break
        if columns is None:
            raise ValueError("Required official master headers were not found in the first 20 rows")
        items = []
        numeric_codes = 0
        for cells in worksheet.iter_rows(min_row=header_row + 1):
            code_cell = cells[columns["code"]]
            if code_cell.value is None:
                if cells[columns["name"]].value is not None:
                    raise ValueError(f"Item name without code at row {cells[0].row}")
                continue  # A report footer is not an item.
            selected = {key: cells[index] for key, index in columns.items()}
            if any(c.data_type == "f" for c in selected.values()):
                raise ValueError(f"Formula in official source fields at row {code_cell.row}")
            item = {key: text(cell.value) for key, cell in selected.items()}
            item["code"] = code_text(code_cell)
            if not item["code"].strip() or not item["name"].strip():
                raise ValueError(f"Missing item code or name at row {code_cell.row}")
            numeric_codes += isinstance(code_cell.value, (int, float))
            items.append(item)
        return items, header_row, columns, numeric_codes
    finally:
        workbook.close()


def convert(source, output_dir, sheet=SHEET):
    items, header_row, columns, numeric_codes = read_items(source, sheet)
    if not items:
        raise ValueError("The official item sheet contains no records")
    structured = {normalize(item["brand"]) for item in items
                  if normalize(item["brand"]) not in GENERIC}
    frequencies = Counter()
    for item in items:
        fragment, _, _ = name_parts(item["name"], structured)
        if plausible_brand(fragment, structured):
            frequencies[normalize(fragment)] += 1
    repeated = {key for key, count in frequencies.items() if count >= 2}
    statuses, evidence_counts = Counter(), Counter()
    for item in items:
        parsed, evidence = parse_item(item["name"], item["brand"], structured, repeated)
        item.update(parsed)
        statuses[item["brandStatus"]] += 1
        evidence_counts[evidence] += 1
    codes = Counter(item["code"] for item in items)
    stats = {
        "schemaVersion": 1,
        "sourceFile": source.name,
        "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sourceSheet": sheet,
        "headerRow": header_row,
        "sourceColumns": {key: index + 1 for key, index in columns.items()},
        "rows": len(items),
        "uniqueCodes": len(codes),
        "duplicateCodeRows": sum(count - 1 for count in codes.values()),
        "leadingZeroCodes": sum(item["code"].startswith("0") for item in items),
        "numericSourceCodes": numeric_codes,
        "structuredBrandRows": sum(bool(normalize(item["brand"])) for item in items),
        "structuredBrandDictionarySize": len(structured),
        "repeatedCandidateDictionarySize": len(repeated - structured),
        "brandStatusCounts": dict(sorted(statuses.items())),
        "modelEvidenceCounts": dict(sorted(evidence_counts.items())),
        "needsBrandConfirmationRows": sum(count for status, count in statuses.items() if status != "explicit"),
        "policy": "brand preserves source Merk; inferred candidates never confirm a brand; candidate and unknown require manual confirmation; generic is reserved for manual confirmation",
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "master_items.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(items)
    stats["csvSha256"] = hashlib.sha256(csv_path.read_bytes()).hexdigest()
    (output_dir / "master_items.stats.json").write_text(
        json.dumps(stats, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return stats


def self_test():
    """Synthetic regression cases only; no business data is embedded in code."""
    from types import SimpleNamespace

    known = {"acme", "north star", "3m"}
    repeated = {"nova"}
    examples = [
        ("Lamp, ACME | A-12", "ACME", "ACME", "explicit", "A-12"),
        ("Lamp, acme #A-12", "", "acme", "candidate", "A-12"),
        ("Lamp, North Star | NS-1", "", "North Star", "candidate", "NS-1"),
        ("Lamp, nova | N-1", "", "nova", "candidate", "N-1"),
        ("Lamp, UNKNOWNCO | X-1", "", "UNKNOWNCO", "unknown", "X-1"),
        ("Lamp, PUTIH | A-12", "", "", "unknown", "A-12"),
        ("Besi Hollow, #0.4 mm", "", "", "unknown", ""),
        ("Lamp, Wall Mounting", "", "", "unknown", ""),
        ("Lamp, LOKAL", "LOKAL", "", "unknown", ""),
        ("Lamp, ACME | A-12", "  ACME  ", "ACME", "explicit", "A-12"),
        ("Tape, 3M #T-1", "", "3M", "candidate", "T-1"),
        ("Lamp, CM | C-1", "CM", "", "explicit", "C-1"),
        ("Lamp, 24 V | 12 mm", "", "", "unknown", ""),
        ("Lamp without comma", "", "", "unknown", ""),
    ]
    for name, raw, candidate, status, model in examples:
        result, _ = parse_item(name, raw, known, repeated)
        assert result == {"brandCandidate": candidate, "brandStatus": status, "model": model}, (name, result)
    assert code_text(SimpleNamespace(value="00123", number_format="General", coordinate="B7")) == "00123"
    assert code_text(SimpleNamespace(value=123, number_format="000000", coordinate="B7")) == "000123"
    assert code_text(SimpleNamespace(value=123.0, number_format="General", coordinate="B7")) == "123"
    print(f"PASS: {len(examples)} parser cases and 3 code preservation cases")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parents[1] / ".local" / "master")
    parser.add_argument("--sheet", default=SHEET)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
    if args.source:
        stats = convert(args.source, args.output_dir, args.sheet)
        print(json.dumps({key: stats[key] for key in ("rows", "uniqueCodes", "leadingZeroCodes", "brandStatusCounts", "modelEvidenceCounts")}, ensure_ascii=False))
        print(f"Saved {args.output_dir / 'master_items.csv'}")
    elif not args.self_test:
        parser.error("source is required unless --self-test is used")


if __name__ == "__main__":
    main()
